
import os
import shutil
from django.conf import settings as django_settings
import gspread
from mailmerge import MailMerge
from docx import Document
from docxcompose.composer import Composer
import pathlib


file =os.path.join(django_settings.STATIC_ROOT)
template = "static/api/voucher.docx"
composed = os.path.join(file, "finals")

files = []

def create(od, do, sheet, request):
    
    try:
        gc = gspread.service_account(filename='static/api/client.json')
        sh = gc.open_by_key('1A41ogDqUWtHUALXmBpjP8Sj_ZRN8mq9c7t43LjYm2x4')
        worksheet = sh.get_worksheet(sheet-1)
    except:
        return -1


    for i in range(od, do+1):
        ime = worksheet.cell(i, 1).value
        prezime = worksheet.cell(i, 2).value
        broj = worksheet.cell(i, 3).value
        gmail =  worksheet.cell(i, 4).value
        datum =  worksheet.cell(i, 5).value
        ispuni(ime, prezime, broj, gmail, datum)

    try:
        spoji()
        ocisti()


   
    except Exception as e:
        print(e)


    
    return 0


def ispuni(ime, prezime, broj, gmail,datum):
    try:
        print("pocetak docx " + ime)
        document = MailMerge(template)
        document.merge(Ime=ime, prezime=prezime, email=gmail, broj = broj, rodenje=datum)
     
        voucher_path = os.path.join(file, "voucheri", "vouchertest" + broj + ".docx")
        print(voucher_path)
   
        document.write(voucher_path)
        print("docx sejvan " + ime)

        

    except Exception as e:
        print(e)



def spoji():
    
    dir = os.path.join(django_settings.STATIC_ROOT, "voucheri")
    for i in pathlib.Path(dir).glob('*.docx'):
        file_path = os.path.join(dir, i)
        print(file_path)
        files.append(file_path)
 
    
    final_path = os.path.join(file, "voucheri", "final.docx")
    result = Document(files[0])
    result.add_page_break()
    composer = Composer(result)

    for i in range(1, len(files)):
        doc = Document(files[i])

        if i != len(files) - 1:
            doc.add_page_break()

        composer.append(doc)
  
    composer.save(final_path)
    print("saved")


def ocisti():
    dir = os.path.join(django_settings.STATIC_ROOT, "voucheri")
    for i in pathlib.Path(dir).glob('*.docx'):
        print(i)
        if i=="/app/static/voucheri/final.docx":
            print("finalan")
            continue
    
        file_path = os.path.join(dir, i)
        print(file_path)
        os.remove(file_path)

